from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from ic_interviewer.ingest import load_resume
from ic_interviewer.ingest import docx_reader, md_reader, pdf_reader


def test_md_reader_loads_text(tmp_path: Path) -> None:
    path = tmp_path / "resume.md"
    path.write_text("# Alice\nRTL engineer\n", encoding="utf-8")

    result = md_reader.load_resume(str(path))

    assert result["source_path"] == str(path)
    assert result["file_type"] == "md"
    assert "RTL engineer" in result["raw_text"]


def test_docx_reader_loads_text(tmp_path: Path) -> None:
    path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Alice Candidate")
    document.add_paragraph("Built interview tooling")
    document.save(path)

    result = docx_reader.load_resume(str(path))

    assert result["file_type"] == "docx"
    assert "Alice Candidate" in result["raw_text"]
    assert "Built interview tooling" in result["raw_text"]


def test_pdf_reader_loads_text(tmp_path: Path) -> None:
    path = tmp_path / "resume.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    with path.open("wb") as fh:
        writer.write(fh)

    result = pdf_reader.load_resume(str(path))

    assert result["source_path"] == str(path)
    assert result["file_type"] == "pdf"
    assert isinstance(result["raw_text"], str)


def test_resume_loader_dispatches_by_extension(tmp_path: Path) -> None:
    path = tmp_path / "resume.md"
    path.write_text("candidate content", encoding="utf-8")

    result = load_resume(str(path))

    assert result["file_type"] == "md"
    assert result["raw_text"] == "candidate content"


def test_missing_file_raises_clear_error(tmp_path: Path) -> None:
    missing = tmp_path / "missing.md"

    with pytest.raises(FileNotFoundError, match="Resume file not found"):
        load_resume(str(missing))


def test_unsupported_file_type_raises_clear_error(tmp_path: Path) -> None:
    path = tmp_path / "resume.txt"
    path.write_text("plain text", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported resume format"):
        load_resume(str(path))
